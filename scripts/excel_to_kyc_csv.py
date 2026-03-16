#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convierte el Excel "Formulario KYC P.F corregido.xlsx" al CSV que usa el módulo
(ghr_compliance/data/kyc_pf_persona_fisica.csv).

Requisito: pip install openpyxl

Uso:
  python excel_to_kyc_csv.py "C:\\Users\\Leand\\Downloads\\trabajo\\Formulario KYC P.F corregido.xlsx"

  o desde la carpeta scripts:
  python excel_to_kyc_csv.py "ruta/al/Formulario KYC P.F corregido.xlsx"

El CSV se escribe en ../data/kyc_pf_persona_fisica.csv (respecto a este script).
"""
from __future__ import print_function

import csv
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    print("Instala openpyxl: pip install openpyxl")
    sys.exit(1)


def slug(s):
    """Convierte etiqueta a nombre técnico x_kyc_xxx."""
    s = (s or "").strip().lower()
    s = re.sub(r"[^\w\s]", "", s)
    s = re.sub(r"\s+", "_", s)
    s = s.strip("_")[:60]
    return "x_kyc_" + s if s else "x_kyc_campo"


def guess_type(label):
    """Inferir tipo de campo por el texto de la etiqueta."""
    lab = (label or "").lower()
    if any(x in lab for x in ["sí", "si", "no?", "¿posee", "¿tiene", "¿es o ha sido", "¿realiza", "¿cuenta con", "¿recibe", "¿tiene algún"]):
        return "Boolean"
    if "fecha" in lab or "date" in lab:
        return "Date"
    if "monto" in lab or "rd$" in lab or "pesos" in lab or "patrimonio" in lab:
        return "Monetary"
    if "número" in lab or "no." in lab or "cantidad" in lab:
        if "teléfono" in lab or "contacto" in lab or "documento" in lab or "pasaporte" in lab or "cuenta" in lab:
            return "Char"
        return "Integer"
    if "meses" in lab or "años" in lab:
        return "Integer"
    if "origen" in lab and "ingresos" in lab:
        return "Selection"
    if "sector" in lab or "formulario" in lab or "estado civil" in lab or "sexo" in lab:
        return "Selection"
    if "tipo" in lab and ("documento" in lab or "cuenta" in lab or "transacción" in lab or "vinculación" in lab):
        return "Selection"
    if "país" in lab or "nacionalidad" in lab:
        return "Many2one (res.country)"
    if "propósito" in lab or "origen de los fondos" in lab or "descripción" in lab:
        return "Text"
    return "Char"


def is_section(text):
    """True si el texto parece un encabezado de sección (ej. '1. Datos...', '1.1 FATCA')."""
    if not text or not isinstance(text, str):
        return False
    t = text.strip()
    return bool(re.match(r"^\d+(\.\d+)?\.\s+", t) or re.match(r"^\d+(\.\d+)?\s+[A-Za-z]", t))


def cell_text(cell):
    if cell is None:
        return ""
    v = getattr(cell, "value", None)
    if v is None:
        return ""
    return str(v).strip()


# Opciones sueltas que no son etiquetas de campo
SKIP_VALUES = frozenset({
    "sí", "si", "no", "privado", "público", "cédula", "pasaporte",
    "residente", "ciudadano", "w8", "w9", "indique cual:", "masculino", "femenino",
})


def extract_fields_from_sheet(ws):
    """Recorre la hoja y devuelve lista de (seccion_actual, nombre_campo, es_seccion)."""
    section = ""
    rows = []
    seen = set()

    for row in ws.iter_rows(max_row=500):
        texts = [cell_text(c) for c in row[:6]]
        line = " ".join(t for t in texts if t).strip()
        if not line:
            continue

        if is_section(line):
            section = line
            rows.append((section, "", True))
            continue

        for part in texts:
            if not part or len(part) < 2:
                continue
            part_lower = part.lower().strip()
            if part_lower in SKIP_VALUES:
                continue
            if is_section(part):
                section = part
                rows.append((section, "", True))
                continue
            if part in seen:
                continue
            if len(part) > 120:
                continue
            seen.add(part)
            rows.append((section, part, False))

    return rows


FIELDNAMES = [
    "Sección",
    "Nombre del Campo",
    "Nombre Técnico del Campo",
    "Tipo de Campo",
    "Optionset Values",
    "Campo en res.partner",
]


def read_excel_as_table(ws):
    """Si la primera fila es cabecera (Sección, Nombre del Campo, ...), devuelve lista de dicts; si no, None."""
    rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
    if len(rows) < 1:
        return None
    first = [str(v).strip() if v is not None else "" for v in rows[0][:6]]
    if not first[0] or "secci" not in first[0].lower():
        return None
    if "nombre" not in (first[1] or "").lower() or "campo" not in (first[1] or "").lower():
        return None
    out = []
    for row in ws.iter_rows(min_row=2, max_row=500, values_only=True):
        cells = list(row)[:6]
        vals = ["" if v is None else str(v).strip() for v in cells]
        seccion, nombre, tecnico, tipo, options, partner = (vals + [""] * 6)[:6]
        if not nombre and not tecnico:
            if seccion:
                out.append({
                    "Sección": seccion,
                    "Nombre del Campo": "",
                    "Nombre Técnico del Campo": "",
                    "Tipo de Campo": "",
                    "Optionset Values": "",
                    "Campo en res.partner": "",
                })
            continue
        if not tecnico and nombre:
            tecnico = slug(nombre)
        if not tipo and nombre:
            tipo = guess_type(nombre)
        out.append({
            "Sección": seccion or (out[-1]["Sección"] if out else ""),
            "Nombre del Campo": nombre,
            "Nombre Técnico del Campo": tecnico,
            "Tipo de Campo": tipo,
            "Optionset Values": options,
            "Campo en res.partner": partner,
        })
    return out


def main():
    if len(sys.argv) < 2:
        print("Uso: python excel_to_kyc_csv.py <ruta_al_excel.xlsx>")
        sys.exit(1)

    xlsx_path = os.path.abspath(sys.argv[1])
    if not os.path.isfile(xlsx_path):
        print("No se encontró el archivo:", xlsx_path)
        sys.exit(1)

    script_dir = os.path.dirname(os.path.abspath(__file__))
    csv_path = os.path.join(script_dir, "..", "data", "kyc_pf_persona_fisica.csv")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    out = read_excel_as_table(ws)
    if out is None or len(out) == 0:
        raw = extract_fields_from_sheet(ws)
        wb.close()
        out = []
        current_section = ""
        for sec, label, is_header in raw:
            if is_header:
                if sec:
                    current_section = sec
                continue
            if not label or not current_section:
                continue
            tech = slug(label)
            tipo = guess_type(label)
            out.append({
                "Sección": current_section,
                "Nombre del Campo": label,
                "Nombre Técnico del Campo": tech,
                "Tipo de Campo": tipo,
                "Optionset Values": "",
                "Campo en res.partner": "",
            })
    else:
        wb.close()

    os.makedirs(os.path.dirname(csv_path), exist_ok=True)
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDNAMES)
        w.writeheader()
        w.writerows(out)

    print("Escrito %d filas en %s" % (len(out), csv_path))


if __name__ == "__main__":
    main()
